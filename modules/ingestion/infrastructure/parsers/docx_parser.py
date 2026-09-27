from __future__ import annotations

from pathlib import Path

from shared.domain.types import Document
from modules.ingestion.infrastructure.parsers.base import Parser, parser_registry


@parser_registry.register("docx_python_docx", "Extracts text/tables from .docx via python-docx.")
class DocxParser(Parser):
    name = "docx_python_docx"

    @classmethod
    def supported_extensions(cls) -> tuple[str, ...]:
        return (".docx",)

    def parse(self, file_path):
        import docx
        path = Path(file_path)
        doc = docx.Document(str(path))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))
        text = "\n".join(parts)
        return Document(source_filename=path.name, text=text, parser_name=self.name,
                         metadata={"num_paragraphs": len(doc.paragraphs), "num_chars": len(text)})
